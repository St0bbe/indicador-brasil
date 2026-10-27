"""Inspeciona planilhas XLSX/CSV no ZIP oficial Sinesp VDE, sem alterar dados.

Uso: python scripts/inspecionar_sinesp.py basededadosvde.zip
"""
import csv
import io
import sys
import zipfile
from pathlib import Path
from openpyxl import load_workbook

if len(sys.argv) != 2:
    raise SystemExit("Uso: python scripts/inspecionar_sinesp.py arquivo.zip")
arquivo = Path(sys.argv[1])
if not zipfile.is_zipfile(arquivo):
    raise SystemExit("O arquivo fornecido nao e ZIP valido")

with zipfile.ZipFile(arquivo) as z:
    membros = [m for m in z.infolist() if not m.is_dir() and
               m.filename.lower().endswith((".xlsx", ".csv"))]
    if not membros:
        raise SystemExit("ZIP nao contem planilhas XLSX ou CSV")
    for membro in membros:
        print("ARQUIVO:", membro.filename, "TAMANHO:", membro.file_size)
        if membro.file_size > 250_000_000:
            print("IGNORADO: tamanho excede limite de seguranca")
            continue
        conteudo = z.read(membro)
        if membro.filename.lower().endswith(".xlsx"):
            wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
            for aba in wb:
                print("ABA:", aba.title, "LINHAS:", aba.max_row, "COLUNAS:", aba.max_column)
                for linha in aba.iter_rows(min_row=1, max_row=min(aba.max_row, 8), values_only=True):
                    print(repr(linha[:20]))
            wb.close()
        else:
            for encoding in ("utf-8-sig", "latin-1"):
                try:
                    texto = conteudo.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            amostra = texto[:10000]
            dialect = csv.Sniffer().sniff(amostra, delimiters=";,\\t|")
            for linha in list(csv.reader(io.StringIO(texto), dialect))[:8]:
                print(repr(linha[:20]))
